"""
run_benchmark.py — end-to-end accuracy benchmark for the container pipeline.

Runs the REAL pipeline (Docling + Surya OCR -> reducer w/ overrides -> DeepSeek
extraction -> DeepSeek access scoring) over a worklist of (filename, brand)
pairs and grades the 12 deliverable parameters + the Access Score against the
PDF-derived truth verdicts (<pdf-stem>__<BRAND>.json) produced by the
project's offline eval harness (not included in this repo).

Usage (see benchmark/README.md):
    python container/benchmark/run_benchmark.py \
        --pdf-dir <dir with the policy PDFs> \
        --worklist <PA_Business_Rules.xlsx | pairs.csv> \
        --truth-dir <directory of truth verdict JSONs> \
        [--out benchmark_report.md] [--runs 1]

Grading reimplements the norm()/struct_match() helpers from the project's
offline eval harness (not included in this repo) so the match semantics are
identical to the validated rescore: exact match after normalization,
leading-number equivalence, NA/No/none "absent" equivalence, yes/required
equivalence, and containment.

IMPORTANT: importing this module performs NO LLM call and NO heavy import —
pipeline.core (which pulls in docling/torch) is imported lazily inside main().
At runtime it needs the DeepSeek env (DEEPSEEK_API_KEY [+ optional gateway
vars], see pipeline/llm.py) and the docling/surya stack installed.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from pathlib import Path

# Make `pipeline` importable when run as a script (benchmark/ lives in container/).
_CONTAINER_DIR = Path(__file__).resolve().parent.parent
if str(_CONTAINER_DIR) not in sys.path:
    sys.path.insert(0, str(_CONTAINER_DIR))

# The 12 deliverable parameters, in deliverable order (matches rescore_full.py).
PARAMS = [
    "Age", "Step Therapy Requirements Documented in Policy",
    "Number of Steps through Brands", "Number of Steps through Generic",
    "Step through-Phototherapy", "TB Test required", "Quantity Limits",
    "Specialist Types", "Initial Authorization Duration(in-months)",
    "Reauthorization Duration(in-months)", "Reauthorization Required",
    "Reauthorization Requirements Documented in Policy",
]

# Structural params graded mechanically with high confidence (rescore_full.py's
# STRUCT set); the free-text params share the same matcher but its containment
# clause makes them a LENIENT grade — flagged as such in the report.
STRUCT = {
    "Age", "Number of Steps through Brands", "Number of Steps through Generic",
    "Step through-Phototherapy", "TB Test required",
    "Initial Authorization Duration(in-months)",
    "Reauthorization Duration(in-months)", "Reauthorization Required",
}

# Short column labels for the per-row matrix table.
PARAM_ABBR = {
    "Age": "Age",
    "Step Therapy Requirements Documented in Policy": "StepTxt",
    "Number of Steps through Brands": "StepsBr",
    "Number of Steps through Generic": "StepsGen",
    "Step through-Phototherapy": "Photo",
    "TB Test required": "TB",
    "Quantity Limits": "QL",
    "Specialist Types": "Spec",
    "Initial Authorization Duration(in-months)": "InitAuth",
    "Reauthorization Duration(in-months)": "ReauthDur",
    "Reauthorization Required": "ReauthReq",
    "Reauthorization Requirements Documented in Policy": "ReauthTxt",
}


# ── Match semantics (reimplemented from the project's offline eval harness,
#    not included in this repo) ───────────────────────────────────────────────
def norm(s) -> str:
    """Lowercase, collapse every non-alphanumeric run to a single space."""
    s = str(s).strip().lower()
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s


def struct_match(new_val, derived) -> bool:
    """rescore_full.py's semantic matcher: exact normalized match, leading-number
    equivalence, absent (NA/No/none/...) equivalence, yes/required equivalence,
    and containment either way."""
    a = norm(new_val)
    b = norm(derived)
    if a == b:
        return True
    # numeric equivalence / contained tokens
    an = re.findall(r"\d+", a)
    bn = re.findall(r"\d+", b)
    if an and bn and an[0] == bn[0]:
        return True
    # NA/No/none equivalence as "absent"
    absent = {"na", "no", "none", "not specified", "unspecified", ""}
    if a in absent and b in absent:
        return True
    # yes/required equivalence
    if a in {"yes", "y", "required"} and b in {"yes", "y", "required"}:
        return True
    if a in b or b in a:
        return True
    return False


def access_match(predicted, derived) -> bool:
    """Bucket agreement: compare the first integer found in each value."""
    pn = re.findall(r"\d+", str(predicted))
    dn = re.findall(r"\d+", str(derived))
    if not pn or not dn:
        return norm(predicted) == norm(derived)
    return pn[0] == dn[0]


# ── Worklist loading (xlsx Submissions tab or filename,brand csv) ─────────────
def load_worklist(path: Path) -> list:
    """Return [(filename, BRAND), ...] from an xlsx (Submissions tab — the
    original PA_Business_Rules.xlsx shape) or a 2-column csv."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl  # local import: only needed for xlsx worklists
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb["Submissions"] if "Submissions" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
        fi, bi = hdr.index("Filename"), hdr.index("Brand")
        pairs = []
        for r in rows[1:]:
            if r[fi] and r[bi]:
                pairs.append((str(r[fi]).strip(), str(r[bi]).strip().upper()))
        return pairs
    # csv: filename,brand — with or without a header row.
    pairs = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            if len(row) < 2 or not row[0].strip():
                continue
            if row[0].strip().lower() == "filename":
                continue  # header
            pairs.append((row[0].strip(), row[1].strip().upper()))
    return pairs


def load_truth(truth_dir: Path, filename: str, brand: str):
    """Load the verdict JSON for (filename, brand); None when absent/unreadable."""
    stem = Path(filename).stem
    vf = truth_dir / f"{stem}__{brand}.json"
    if not vf.exists():
        return None
    try:
        return json.loads(vf.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


# ── Report writing ────────────────────────────────────────────────────────────
def write_report(out_path: Path, rows: list, param_counts: dict, access_counts: dict,
                 skipped: list, errors: list, elapsed_s: float) -> None:
    """rows: [{filename, brand, matches: {param: bool}, access_ok: bool,
               predicted_access, truth_access}, ...]"""
    lines = []
    lines.append("# Container pipeline benchmark report")
    lines.append("")
    lines.append(f"Rows graded: {len(rows)} | skipped (no truth verdict): "
                 f"{len(skipped)} | pipeline errors: {len(errors)} | "
                 f"wall time: {elapsed_s:.0f}s")
    lines.append("")

    # Per-parameter accuracy matrix.
    lines.append("## Per-parameter accuracy")
    lines.append("")
    lines.append("| Parameter | Grade | Accuracy | OK / Total |")
    lines.append("|---|---|---|---|")
    for p in PARAMS:
        ok, total = param_counts[p]["ok"], param_counts[p]["total"]
        pct = f"{100 * ok / total:.1f}%" if total else "n/a"
        grade = "structural" if p in STRUCT else "lenient (containment)"
        lines.append(f"| {p} | {grade} | {pct} | {ok} / {total} |")
    all_ok = sum(param_counts[p]["ok"] for p in PARAMS)
    all_total = sum(param_counts[p]["total"] for p in PARAMS)
    if all_total:
        lines.append(f"| **All parameters** | | **{100 * all_ok / all_total:.1f}%** "
                     f"| {all_ok} / {all_total} |")
    lines.append("")

    # Access Score agreement.
    lines.append("## Access Score agreement")
    lines.append("")
    a_ok, a_total = access_counts["ok"], access_counts["total"]
    pct = f"{100 * a_ok / a_total:.1f}%" if a_total else "n/a"
    lines.append(f"Bucket agreement with the truth verdicts: **{pct}** "
                 f"({a_ok} / {a_total})")
    lines.append("")

    # Per-row matrix.
    lines.append("## Per-row results")
    lines.append("")
    abbrs = [PARAM_ABBR[p] for p in PARAMS]
    lines.append("| Filename | Brand | " + " | ".join(abbrs) + " | Params OK | Access (pred/truth) |")
    lines.append("|---" * (len(abbrs) + 4) + "|")
    for r in rows:
        marks = ["Y" if r["matches"][p] else "N" for p in PARAMS]
        n_ok = sum(1 for p in PARAMS if r["matches"][p])
        acc = f"{r['predicted_access']}/{r['truth_access']}" + ("" if r["access_ok"] else " (X)")
        lines.append(f"| {r['filename']} | {r['brand']} | " + " | ".join(marks) +
                     f" | {n_ok}/12 | {acc} |")
    lines.append("")

    if skipped:
        lines.append("## Skipped (no truth verdict found)")
        lines.append("")
        for fn, br in skipped:
            lines.append(f"- {fn} | {br}")
        lines.append("")
    if errors:
        lines.append("## Pipeline errors")
        lines.append("")
        for fn, br, err in errors:
            lines.append(f"- {fn} | {br}: {err}")
        lines.append("")

    out_path.write_text("\n".join(lines), encoding="utf-8")


# ── Main ──────────────────────────────────────────────────────────────────────
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Benchmark the container pipeline "
                                 "against the eval truth verdicts.")
    ap.add_argument("--pdf-dir", required=True, help="directory holding the policy PDFs")
    ap.add_argument("--worklist", required=True,
                    help="xlsx (Submissions tab, e.g. PA_Business_Rules.xlsx) or "
                         "csv of filename,brand pairs")
    ap.add_argument("--truth-dir", required=True,
                    help="directory of truth verdict JSONs (<pdf-stem>__<BRAND>.json)")
    ap.add_argument("--out", default="benchmark_report.md", help="output report path (markdown)")
    ap.add_argument("--runs", type=int, default=1,
                    help="confidence runs per (pdf, brand) pair (1..5, default 1)")
    args = ap.parse_args(argv)

    import os
    if not os.environ.get("DEEPSEEK_API_KEY"):
        print("ERROR: DEEPSEEK_API_KEY is not set — the pipeline cannot call the "
              "model. See pipeline/llm.py for the full env contract.", file=sys.stderr)
        return 2

    # Heavy import deferred to here (docling/torch via pipeline.core) so that
    # importing this module stays side-effect free.
    from pipeline import core

    pdf_dir = Path(args.pdf_dir)
    truth_dir = Path(args.truth_dir)
    worklist = load_worklist(Path(args.worklist))
    print(f"Worklist: {len(worklist)} (filename, brand) pairs")

    rows, skipped, errors = [], [], []
    param_counts = {p: {"ok": 0, "total": 0} for p in PARAMS}
    access_counts = {"ok": 0, "total": 0}
    t_start = time.time()

    for i, (filename, brand) in enumerate(worklist, 1):
        print(f"[{i}/{len(worklist)}] {filename} | {brand}")
        truth = load_truth(truth_dir, filename, brand)
        if truth is None:
            print("  no truth verdict — skipped")
            skipped.append((filename, brand))
            continue
        pdf_path = pdf_dir / filename
        if not pdf_path.exists():
            print("  PDF not found — recorded as error")
            errors.append((filename, brand, "PDF not found"))
            continue

        try:
            # Spans/consensus off: this benchmark grades the 13 deliverables only.
            out = core.extract_one(pdf_path.read_bytes(), filename, brand,
                                   runs=args.runs, explain_spans=False,
                                   consensus=False)
        except Exception as e:  # noqa: BLE001 — keep benchmarking the rest
            print(f"  PIPELINE ERROR: {e}")
            errors.append((filename, brand, str(e)))
            continue

        result = out["result"]
        pmap = {p.get("param"): p for p in truth.get("parameters", [])}
        matches = {}
        for p in PARAMS:
            derived = str(pmap.get(p, {}).get("derived", ""))
            ok = struct_match(result.get(p, ""), derived)
            matches[p] = ok
            param_counts[p]["ok"] += int(ok)
            param_counts[p]["total"] += 1

        truth_access = str(truth.get("access_score", {}).get("derived", ""))
        predicted_access = str(result.get("Access Score", ""))
        a_ok = access_match(predicted_access, truth_access)
        access_counts["ok"] += int(a_ok)
        access_counts["total"] += 1

        n_ok = sum(matches.values())
        print(f"  params {n_ok}/12 | access {predicted_access} vs {truth_access} "
              f"({'ok' if a_ok else 'MISMATCH'})")
        rows.append({"filename": filename, "brand": brand, "matches": matches,
                     "access_ok": a_ok, "predicted_access": predicted_access,
                     "truth_access": truth_access})

    write_report(Path(args.out), rows, param_counts, access_counts,
                 skipped, errors, time.time() - t_start)
    print(f"Report written to {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
